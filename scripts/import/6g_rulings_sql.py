"""Phase 6g: turn the owner-accepted review workbook into INSERTs for owner_station_rulings.

Every identity string is read from the workbook and written out verbatim (never retyped). The output ends with
the md5 the loaded rows must reproduce in the database, so a transcription error cannot pass unnoticed.
Usage: python3 scripts/import/6g_rulings_sql.py deliverables/phase-6c-zero-station-regions-review.xlsx > out.sql
"""
import hashlib, re, sys
import openpyxl

path = sys.argv[1]
sha = hashlib.sha256(open(path, 'rb').read()).hexdigest()
ws = openpyxl.load_workbook(path)['Names to review']
header = [c.value for c in ws[1]]
assert header[:8] == ['Region', 'Source name (as in files)', 'Other spellings', 'Found in', 'Rows',
                      'Proposed Station', 'Proposed Unit', 'Why proposed'], header

def q(v):
    return 'NULL' if v is None or str(v).strip() == '' else "'" + str(v).replace("'", "''") + "'"

rows = []
held = []
for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    region, src, other, found, _n, st, un, why = r[:8]
    if not src:
        continue
    if r[8] or r[9]:
        raise SystemExit(f'row {i}: owner columns are filled; this generator is for "accept the proposals" only')
    # The trailing-number rule misreads a number that is part of the name ("الكيلو 21- 1", "ك/1/21",
    # "1&2"), leaving a Station name ending in a separator. Such rows are HELD for the owner, never created.
    if re.search(r'[-/&]$', str(st).strip()) or (un and re.search(r'[-/&] \d+$', str(un))):
        held.append((i, region, src))
        continue
    evidence = f'review workbook row {i}; found in: {found}'
    rows.append((region, src, other, st, un, why, evidence))

print(f'-- source workbook sha256 {sha}')
print('BEGIN;')
print("INSERT INTO owner_station_rulings (ruling_set, region_id, source_name_raw, other_spelling_raw, station_name, unit_name, proposal_reason, evidence)")
print('SELECT \'6G-2026-09-24\', g.id, v.src, v.oth, v.st, v.un, v.why, v.ev FROM (VALUES')
print(',\n'.join(f'  ({q(reg)}, {q(s)}, {q(o)}, {q(t)}, {q(u)}, {q(w)}, {q(e)})' for reg, s, o, t, u, w, e in rows))
print(') v(region, src, oth, st, un, why, ev) JOIN regions g ON g.name = v.region;')
print('COMMIT;')
digest = hashlib.md5('\n'.join('|'.join('' if x is None else str(x) for x in (reg, s, o, t, u)) for reg, s, o, t, u, _w, _e in sorted(rows, key=lambda x: (x[0], x[1]))).encode()).hexdigest()
print(f'-- rows {len(rows)} md5 {digest}')
print(f'-- held {len(held)}: workbook rows ' + ', '.join(str(h[0]) for h in held))
