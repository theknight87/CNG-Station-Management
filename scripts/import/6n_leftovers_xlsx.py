"""Owner review file for what Phase 6n did not change. Run in the 6n work dir after 6n_snapshot_update.py (leftovers)."""
import json, re, sys, unicodedata
import openpyxl
from openpyxl.styles import Font

def n(s): return re.sub(r'\s+', ' ', unicodedata.normalize('NFKC', '' if s is None else str(s))).strip().lower()
F = {r['row']: r for r in json.load(open('file_norm.json'))}
S = {d[0]: d for d in ((l.split('\t') + [''] * 8)[:8] for l in open('system_srvs.tsv', encoding='utf-8').read().splitlines())}
L = json.load(open('6n_leftovers.json'))
fr = lambda r: [r['area'], r['station_txt'], r['location_txt'], r['serial_norm']['raw'], r['pressure_norm']['raw'],
                r['last_norm']['value'] or r['last_norm']['raw'], r['next_norm']['value'] or r['next_norm']['raw'], r['row']]
sr = lambda d: [d[1], d[2], d[3], d[4] or None, d[5] or None, d[7] or None, d[0]]
key = lambda x: (n(x[1]), n(x[2]), n(x[4]))

o = openpyxl.Workbook(); ws = o.active; ws.title = 'ملخص'; ws.sheet_view.rightToLeft = True
for line in [('البيان', 'العدد'), ('مش واضح مين يقابل مين: في الملف', len(L['amb_f'])), ('مش واضح مين يقابل مين: في النظام', len(L['amb_s'])),
             ('في الملف بس (جديدة؟)', len(L['file_only'])), ('في النظام بس (اتشالت؟)', len(L['sys_only']))]:
    ws.append(line)
ws['A1'].font = ws['B1'].font = Font(bold=True); ws.column_dimensions['A'].width = 40

FH = ['المنطقة', 'المحطة', 'المكان', 'السيريال', 'الضغط', 'آخر معايرة', 'المعايرة الجاية', 'صف الملف']
SH = ['المنطقة', 'المحطة', 'المكان', 'السيريال', 'الضغط', 'Unit', 'system_id']
def sheet(title, hdr, rows):
    d = o.create_sheet(title); d.sheet_view.rightToLeft = True; d.append(hdr)
    for c in d[1]: c.font = Font(bold=True)
    for r in sorted(rows, key=key): d.append(r)
    for i, w in enumerate([9, 30, 9, 18, 12, 14, 14, 38, 38]): d.column_dimensions[chr(65 + i)].width = w
    d.freeze_panes = 'A2'; d.auto_filter.ref = d.dimensions

amb = [['الملف'] + fr(F[i]) for i in L['amb_f']] + [['النظام'] + sr(S[i])[:5] + [None, None, S[i][0]] for i in L['amb_s']]
d = o.create_sheet('مش واضح'); d.sheet_view.rightToLeft = True
d.append(['المصدر'] + FH[:7] + ['صف الملف / system_id'])
for c in d[1]: c.font = Font(bold=True)
for r in sorted(amb, key=lambda x: (n(x[1]), n(x[2]), n(x[3]), n(x[5]), x[0])): d.append(r)
for i, w in enumerate([8, 9, 30, 9, 18, 12, 14, 14, 38]): d.column_dimensions[chr(65 + i)].width = w
d.freeze_panes = 'A2'; d.auto_filter.ref = d.dimensions
sheet('في الملف بس', FH, [fr(F[i]) for i in L['file_only']])
sheet('في النظام بس', SH, [sr(S[i]) for i in L['sys_only']])
o.save(sys.argv[1]); print('ok')
