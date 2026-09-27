"""Owner review file: installed SRVs awaiting their Unit (needs_unit_mapping) or Station (needs_station_mapping).

Usage: python3 unit_station_review_xlsx.py <execute_sql result file> <out.xlsx>
Suggestions are text-similarity hints for the owner only; nothing is decided here (CLAUDE.md §4, §8).
"""
import collections, difflib, json, re, sys, unicodedata
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

o = json.load(open(sys.argv[1]))['result']
j = json.loads(o[o.index('[{'):o.rindex('}]') + 2])[0]['j']

def fold(s):
    s = unicodedata.normalize('NFKC', s or '').replace('ـ', '')
    for a, b in (('أ', 'ا'), ('إ', 'ا'), ('آ', 'ا'), ('ى', 'ي'), ('ة', 'ه')):
        s = s.replace(a, b)
    return re.sub(r'\s+', ' ', s).strip()

fill = PatternFill('solid', fgColor='FFF2CC')
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'ملخص'; ws.sheet_view.rightToLeft = True

# --- needs Unit: one row per valve, dropdown of the Station's Units, suggestion when the raw name carries the Unit number
u = wb.create_sheet('محتاجة Unit'); u.sheet_view.rightToLeft = True
u.append(['المنطقة', 'المحطة', 'الاسم في الملف', 'المكان', 'السيريال', 'الضغط', 'الـ Units في المحطة', 'اقتراح', 'الـ Unit (اختيار)', 'ملاحظات', 'srv_id'])
sug_u = 0
for r in j['unit']:
    units = r['units'] or []
    m = re.search(r'(\d+)\s*$', fold(r['raw']))
    s = None
    if m:
        c = [x for x in units if re.search(r'(^|\D)' + m.group(1) + r'\s*$', fold(x))]
        s = c[0] if len(c) == 1 else None
    sug_u += bool(s)
    u.append([r['region'], r['station'], r['raw'], r['location'], r['serial'], r['pressure'], ' | '.join(units), s, None, None, r['id']])
    row = u.max_row; u.cell(row, 9).fill = fill
    f = '"' + ','.join(x.replace(',', ' ').replace('"', "'") for x in units) + '"'
    if units and len(f) < 250:
        dv = DataValidation(type='list', formula1=f, allow_blank=True, showErrorMessage=False); u.add_data_validation(dv); dv.add(u.cell(row, 9))

# --- needs Station: one row per (Region, raw name), with the closest Station names of that Region as hints
g = collections.OrderedDict()
for r in j['station']:
    g.setdefault((r['region'], r['raw']), []).append(r)
s = wb.create_sheet('محتاجة محطة'); s.sheet_view.rightToLeft = True
s.append(['المنطقة', 'الاسم في الملف', 'عدد الريليفات', 'Stage', 'Storage', 'أقرب أسماء في نفس المنطقة', 'المحطة الصح', 'الـ Unit (لو معروفة)', 'ملاحظات'])
for (reg, raw), rows in g.items():
    names = j['stations'].get(reg, [])
    folded = {fold(n): n for n in names}
    tok = lambda x: {t for t in re.split(r'[\s/()\-&]+', fold(x).replace('ال', '')) if len(t) > 1 and not t.isdigit()}
    rt = tok(raw)
    by_tok = [n for n in names if rt and rt <= tok(n)] + [n for n in names if rt and tok(n) and tok(n) <= rt]
    close = list(dict.fromkeys(by_tok + [folded[x] for x in difflib.get_close_matches(fold(raw), list(folded), n=3, cutoff=0.5)]))[:4]
    s.append([reg, raw, len(rows), sum(x['location'] == 'Stage' for x in rows), sum(x['location'] == 'Storage' for x in rows),
              ' | '.join(close) or None, None, None, None])
    s.cell(s.max_row, 7).fill = fill; s.cell(s.max_row, 8).fill = fill
d = wb.create_sheet('ريليفات من غير محطة'); d.sheet_view.rightToLeft = True
d.append(['المنطقة', 'الاسم في الملف', 'المكان', 'السيريال', 'الضغط', 'srv_id'])
for r in j['station']:
    d.append([r['region'], r['raw'], r['location'], r['serial'], r['pressure'], r['id']])

for sh, widths in ((u, [8, 22, 24, 8, 16, 11, 40, 20, 22, 20, 38]), (s, [8, 30, 10, 8, 8, 60, 28, 22, 24]), (d, [8, 30, 8, 16, 11, 38])):
    for c in sh[1]: c.font = Font(bold=True)
    for i, w in enumerate(widths): sh.column_dimensions[chr(65 + i)].width = w
    sh.freeze_panes = 'A2'; sh.auto_filter.ref = sh.dimensions

for line in [('البيان', 'العدد'), ('ريليفات محتاجة Unit', len(j['unit'])), ('  منها ليها اقتراح من رقم في الاسم', sug_u),
             ('ريليفات محتاجة محطة', len(j['station'])), ('  عدد الأسماء المختلفة', len(g))]:
    ws.append(line)
ws['A1'].font = ws['B1'].font = Font(bold=True); ws.column_dimensions['A'].width = 38
ws['A8'] = 'شيت "محتاجة Unit": اختار الـ Unit لكل ريليف (الاقتراح جاي من الرقم اللي في الاسم، راجعه).'
ws['A9'] = 'شيت "محتاجة محطة": صف لكل اسم؛ اكتب المحطة الصح (أو "جديدة" لو المحطة مش موجودة) والـ Unit لو تعرفها.'
wb.save(sys.argv[2]); print(len(j['unit']), sug_u, len(j['station']), len(g))
