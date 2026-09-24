"""Review workbook: installed-SRV snapshot (رصيد المحطات) vs the SRVs in the system.

Usage: python3 srv_snapshot_review_xlsx.py <snapshot.xlsx> <system.tsv> <out.xlsx>
system.tsv columns: id, region, station, location, serial, set_pressure, mapping_status, unit.
Matching is exact after NFKC/whitespace/case folding only; nothing is decided here - every difference is a proposal
for the owner (CLAUDE.md §6.8, §6.16).
"""
import collections, re, sys, unicodedata
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

def n(s):
    if isinstance(s, float) and s.is_integer():
        s = int(s)
    return re.sub(r'\s+', ' ', unicodedata.normalize('NFKC', '' if s is None else str(s))).strip().lower()

def main(src, tsv, out):
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    rows = list(wb['رصيد المحطات'].iter_rows(values_only=True))
    file_rows = [(i + 1, r) for i, r in enumerate(rows) if i >= 5 and any(r[:6])]
    sysr = [(l.split('\t') + [''] * 8)[:8] for l in open(tsv, encoding='utf-8').read().splitlines() if l.strip()]

    k4 = lambda reg, st, loc, sn: (n(reg), n(st), n(loc), n(sn))
    F = collections.defaultdict(list); S = collections.defaultdict(list)
    for no, r in file_rows: F[k4(r[0], r[1], r[2], r[5])].append((no, r))
    for d in sysr: S[k4(d[1], d[2], d[3], d[4])].append(d)

    out_rows = []  # (type, region, station, location, serial, p_file, p_sys, file_row, sys_id, status, unit, action)
    same = 0
    for k in sorted(set(F) | set(S)):
        fl = list(F.get(k, [])); sl = list(S.get(k, []))
        # exact (incl. pressure) pairs first
        for f in list(fl):
            m = next((s for s in sl if n(s[5]) == n(f[1][3])), None)
            if m: same += 1; fl.remove(f); sl.remove(m)
        # same valve, pressure differs
        while fl and sl:
            (no, r), d = fl.pop(0), sl.pop(0)
            out_rows.append(('ضغط مختلف', r[0], r[1], r[2], r[5], r[3], d[5], no, d[0], d[6], d[7], 'تعديل الضغط'))
        for no, r in fl:
            out_rows.append(('في الملف بس', r[0], r[1], r[2], r[5], r[3], None, no, None, None, None, 'إضافة'))
        for d in sl:
            out_rows.append(('في النظام بس', d[1], d[2], d[3], d[4] or None, None, d[5], None, d[0], d[6], d[7], 'أرشفة'))

    out_rows.sort(key=lambda x: (n(x[1]), n(x[2]), x[0], n(x[3]), n(x[4])))
    o = openpyxl.Workbook(); ws = o.active; ws.title = 'ملخص'; ws.sheet_view.rightToLeft = True
    c = collections.Counter(x[0] for x in out_rows)
    for line in [('بيان', 'العدد'), ('ريليفات في الملف', len(file_rows)), ('ريليفات في النظام', len(sysr)),
                 ('متطابقة (مش محتاجة حاجة)', same), ('ضغط مختلف', c['ضغط مختلف']),
                 ('في الملف بس', c['في الملف بس']), ('في النظام بس', c['في النظام بس']),
                 ('محطات فيها فروق', len({(n(x[1]), n(x[2])) for x in out_rows}))]:
        ws.append(line)
    ws['A10'] = 'اكتب قرارك في عمود "قرار المالك" في شيت الفروق: موافق / رفض / أو اكتب البديل. لو سبته فاضي هيتساب زي ما هو.'
    ws.column_dimensions['A'].width = 34; ws['A1'].font = ws['B1'].font = Font(bold=True)

    d = o.create_sheet('الفروق'); d.sheet_view.rightToLeft = True
    hdr = ['النوع', 'المنطقة', 'المحطة', 'المكان', 'السيريال', 'الضغط في الملف', 'الضغط في النظام', 'صف الملف',
           'Unit في النظام', 'حالة الربط', 'المقترح', 'قرار المالك', 'ملاحظات', 'system_id']
    d.append(hdr)
    for x in out_rows:
        d.append([x[0], x[1], x[2], x[3], x[4], x[5], x[6], x[7], x[10], x[9], x[11], None, None, x[8]])
    fill = PatternFill('solid', fgColor='FFF2CC')
    for cell in d[1]: cell.font = Font(bold=True)
    for r in range(2, d.max_row + 1): d.cell(r, 12).fill = fill
    dv = DataValidation(type='list', formula1='"موافق,رفض"', allow_blank=True, showErrorMessage=False)
    d.add_data_validation(dv); dv.add(f'L2:L{d.max_row}')
    for col, w in zip('ABCDEFGHIJKLMN', [13, 9, 30, 9, 16, 13, 13, 9, 24, 22, 13, 13, 24, 38]):
        d.column_dimensions[col].width = w
    d.freeze_panes = 'A2'; d.auto_filter.ref = d.dimensions
    o.save(out)
    print(len(file_rows), len(sysr), same, dict(c))

if __name__ == '__main__':
    main(*sys.argv[1:4])
