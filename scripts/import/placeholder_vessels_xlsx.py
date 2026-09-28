"""Owner review file: the temporary placeholder storage vessels created by owner ruling 6r.

Usage: python3 placeholder_vessels_xlsx.py <execute_sql result file> <out.xlsx>
Lists each placeholder with its Station/Unit and the SRVs it carries, with blank columns for the real vessel's data.
Nothing is decided here; the owner fills in the real vessels from the site.
"""
import json, sys
import openpyxl
from openpyxl.styles import Font, PatternFill

o = json.load(open(sys.argv[1]))['result'] if sys.argv[1].endswith('.json') else open(sys.argv[1]).read()
if not isinstance(o, str): o = json.dumps(o)
s = o.replace('\\"', '"') if '\\"j\\"' in o else o
rows = json.loads(s[s.index('[{"j":'):s.rindex('}]') + 2])[0]['j']
fill = PatternFill('solid', fgColor='FFF2CC')
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'ملخص'; ws.sheet_view.rightToLeft = True
v = wb.create_sheet('الخزانات المؤقتة'); v.sheet_view.rightToLeft = True
v.append(['المنطقة', 'المحطة', 'الـ Unit', 'الكومبريسور', 'خزانات حقيقية في المحطة', 'عدد الريليفات عليه', 'سيريالات الريليفات',
          'الشركة المصنعة (الحقيقي)', 'الموديل / النوع', 'السيريال', 'السعة', 'ضغط التشغيل', 'ملاحظات', 'vessel_id'])
d = wb.create_sheet('ريليفات الخزانات المؤقتة'); d.sheet_view.rightToLeft = True
d.append(['المنطقة', 'المحطة', 'الـ Unit', 'سيريال الريليف', 'الشركة', 'الضغط', 'srv_id', 'vessel_id'])
n = 0
for r in rows:
    srvs = r['srvs'] or []; n += len(srvs)
    v.append([r['region'], r['station'], r['unit'], r['comp'], r['real_vessels_at_station'], len(srvs),
              ' | '.join(x[0] or '(بدون)' for x in srvs), None, None, None, None, None, None, r['id']])
    for c in range(8, 14): v.cell(v.max_row, c).fill = fill
    for x in srvs: d.append([r['region'], r['station'], r['unit'], x[0], x[1], x[2], x[3], r['id']])
for sh, w in ((v, [8, 24, 24, 20, 12, 10, 45, 20, 18, 18, 10, 12, 24, 38]), (d, [8, 24, 24, 18, 14, 12, 38, 38])):
    for c in sh[1]: c.font = Font(bold=True)
    for i, x in enumerate(w): sh.column_dimensions[chr(65 + i)].width = x
    sh.freeze_panes = 'A2'; sh.auto_filter.ref = sh.dimensions
for line in [('البيان', 'العدد'), ('خزانات مؤقتة', len(rows)), ('ريليفات مربوطة عليها', n),
             ('خزانات مؤقتة من غير ريليفات', sum(not r['srvs'] for r in rows))]:
    ws.append(line)
ws['A1'].font = ws['B1'].font = Font(bold=True); ws.column_dimensions['A'].width = 40
ws['A7'] = 'دي خزانات اتعملت مؤقتًا (مرحلة 6r) عشان ريليفات الـ Storage تتربط بحاجة لحد ما الخزان الحقيقي يتسجل.'
ws['A8'] = 'املا الأعمدة الصفرا ببيانات الخزان الحقيقي من الموقع؛ اللي تسيبه فاضي يفضل فاضي (مفيش بيانات بتتألف).'
wb.save(sys.argv[2]); print(len(rows), n)
