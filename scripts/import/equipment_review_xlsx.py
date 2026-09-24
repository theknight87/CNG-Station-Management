"""Owner review file: installed SRVs awaiting their equipment parent (needs_equipment_mapping).

Usage: python3 equipment_review_xlsx.py <execute_sql result file> <out.xlsx>
One row per valve with the Unit's candidate compressors / storage vessels; the owner writes the chosen equipment.
Nothing is decided here (CLAUDE.md §4: equipment parentage is an explicit human decision).
"""
import collections, json, sys
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

o = json.load(open(sys.argv[1]))['result']
rows = json.loads(o[o.index("[{"):o.rindex("}]") + 2])[0]["json_agg"]

def reason(r):
    kind = 'comps' if r['location'] == 'Stage' else 'vessels'
    n = len(r[kind] or [])
    if r['location'] not in ('Stage', 'Storage'): return 'المكان مش محدد'
    if n == 0: return 'الـ Unit ملهاش ' + ('كمبروسر' if kind == 'comps' else 'خزان') + ' متسجل'
    return f'الـ Unit فيها {n} ' + ('كمبروسر' if kind == 'comps' else 'خزان')

wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'ملخص'; ws.sheet_view.rightToLeft = True
c = collections.Counter(reason(r) for r in rows)
ws.append(['السبب', 'عدد الريليفات']); [ws.append([k, v]) for k, v in c.most_common()]; ws.append(['الإجمالي', len(rows)])
ws['A1'].font = ws['B1'].font = Font(bold=True); ws.column_dimensions['A'].width = 34
ws['A12'] = 'في شيت "الريليفات": اختار المعدة من عمود "المعدة (اختيار)" أو اكتبها في الملاحظات. الفاضي هيفضل زي ما هو.'

d = wb.create_sheet('الريليفات'); d.sheet_view.rightToLeft = True
hdr = ['المنطقة', 'المحطة', 'الـ Unit', 'المكان', 'السيريال', 'الضغط', 'الشركة', 'المعايرة الجاية', 'السبب',
       'المعدات المتاحة في الـ Unit', 'المعدة (اختيار)', 'ملاحظات', 'srv_id']
d.append(hdr); [setattr(x, 'font', Font(bold=True)) for x in d[1]]
fill = PatternFill('solid', fgColor='FFF2CC')
for r in rows:
    cand = (r['comps'] if r['location'] == 'Stage' else r['vessels']) or []
    labels = [f"{i + 1}) {e['label'] or 'بدون سيريال'}" for i, e in enumerate(cand)]
    d.append([r['region'], r['station'], r['unit'], r['location'], r['serial'], r['pressure'], r['manufacturer'],
              r['next_cal'] or r['next_raw'], reason(r), ' | '.join(labels) or None, None, None, r['id']])
    row = d.max_row; d.cell(row, 11).fill = fill
    if labels:
        dv = DataValidation(type='list', formula1='"' + ','.join(l.replace(',', ' ').replace('"', "'")[:60] for l in labels) + '"',
                            allow_blank=True, showErrorMessage=False)
        if len(dv.formula1) < 250: d.add_data_validation(dv); dv.add(d.cell(row, 11))
for col, w in zip('ABCDEFGHIJKLM', [8, 28, 26, 8, 16, 11, 11, 13, 24, 48, 30, 24, 38]): d.column_dimensions[col].width = w
d.freeze_panes = 'A2'; d.auto_filter.ref = d.dimensions
wb.save(sys.argv[2]); print(len(rows), dict(c))
