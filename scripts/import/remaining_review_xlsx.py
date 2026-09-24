"""Build the owner review workbook for the Station names still unlinked after phase 6i.

Rows are the 6i review rows that were either held (suggestion named a different place) or had no suggestion. Each
gets the three closest existing Stations in its Region as HINTS only; the owner writes the Station (existing or new)
and the Unit. Usage: python3 scripts/import/remaining_review_xlsx.py prev_review.xlsx stations.json out.xlsx
"""
import difflib, json, re, sys
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HELD = {13: 'Suggested "بريما / السادات 3" - a different number', 19: 'شبين الكوم has five Stations',
        45: 'East has both نفق العبور and العبور المنطقة الصناعية', 65: 'Suggested "موبيل المعادي" - a different place',
        72: 'Suggested "البراجيل الجديدة" - old vs new', 105: 'West has a separate Station "فويل اب الدائرى"',
        107: 'West has a separate Station "فويل اب الدائرى"', 108: 'West has a separate Station "فويل اب الدائرى"'}

def fold(s):
    s = s.replace('ـ', '')
    s = re.sub('[أإآ]', 'ا', s).replace('ى', 'ي').replace('ة', 'ه')
    return re.sub(r'\s+', '', s)

stations = json.load(open(sys.argv[2]))
by_region = {}
for region, st, units in stations:
    by_region.setdefault(region, []).append((st, units))

prev = openpyxl.load_workbook(sys.argv[1])['Names to review']
wb = Workbook()
ws = wb.active
ws.title = 'Names to review'
ws.sheet_view.rightToLeft = True
ws.append(['Region', 'Source name (as in files)', 'Storage vessels', 'Recovery tanks', 'Gas detectors', 'Hoses',
           'Installed SRVs', 'Total records', 'Why it is here', 'Closest Station 1', 'Closest Station 2', 'Closest Station 3',
           'Your Station name (existing or NEW)', 'Your Unit name (leave empty if unknown)', 'Comment'])
yellow = PatternFill('solid', fgColor='FFF2A8')
for c in ws[1]:
    c.font = Font(bold=True)
    c.alignment = Alignment(wrap_text=True, vertical='top')
n = total = 0
for i, r in enumerate(prev.iter_rows(min_row=2, values_only=True), start=2):
    if i not in HELD and r[8]:
        continue
    region, name = r[0], r[1]
    why = HELD.get(i, 'No similar Station name in this Region - new Station?')
    f = fold(name)
    close = sorted(by_region.get(region, []), key=lambda c: -difflib.SequenceMatcher(None, f, fold(c[0])).ratio())[:3]
    hints = [c[0] + (f"  (Units: {' ، '.join(c[1])})" if c[1] else '') for c in close]
    ws.append([region, name, *r[2:8], why, *hints, None, None, None])
    for col in (13, 14):
        ws.cell(ws.max_row, col).fill = yellow
    n += 1; total += r[7]
for i, w in enumerate([8, 34, 9, 9, 9, 7, 9, 9, 40, 36, 36, 36, 32, 30, 24], start=1):
    ws.column_dimensions[ws.cell(1, i).column_letter].width = w
ws.freeze_panes = 'C2'
ws.auto_filter.ref = ws.dimensions

ref = wb.create_sheet('Existing Stations and Units')
ref.sheet_view.rightToLeft = True
ref.append(['Region', 'Station', 'Units'])
for region, st, units in stations:
    ref.append([region, st, ' ، '.join(units)])
for col, w in zip('ABC', (8, 36, 70)):
    ref.column_dimensions[col].width = w

how = wb.create_sheet('How to fill')
for line in [
    'These names match no Station yet. The three "Closest Station" columns are spelling hints only, not suggestions.',
    'Your Station name: copy an existing Station exactly from the "Existing Stations and Units" sheet, or write a NEW Station name (it will be created in that Region).',
    'Your Unit name: only if the source name says which Unit (e.g. "X 1"); leave empty otherwise. A new Unit is created if needed.',
    'Leave the row empty to keep those records unlinked.',
]:
    how.append([line])
how.column_dimensions['A'].width = 140
wb.save(sys.argv[3])
print(n, 'names,', total, 'records')
