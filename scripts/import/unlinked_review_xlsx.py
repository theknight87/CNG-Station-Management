"""Build the owner review workbook for records whose source Station name matches no Station or ruling.

Inputs are JSON exports from production (names with per-family counts; Stations with their Units), each verified
against the md5 the database reported. Suggestions are REVIEW AIDS only: an exact Unit-name or Station-name match
after spelling folding, else the closest Station spelling (marked "check"). Nothing here is applied automatically.
Usage: python3 scripts/import/unlinked_review_xlsx.py names.json stations.json out.xlsx
"""
import difflib, json, re, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

names = json.load(open(sys.argv[1]))
stations = json.load(open(sys.argv[2]))

def fold(s):
    s = s.replace('ـ', '')
    s = re.sub('[أإآ]', 'ا', s).replace('ى', 'ي').replace('ة', 'ه')
    s = s.translate(str.maketrans('٠١٢٣٤٥٦٧٨٩', '0123456789'))
    return re.sub(r'\s+', '', s)

by_region = {}
for region, st, units in stations:
    by_region.setdefault(region, []).append((st, units))

def suggest(region, name):
    f = fold(name)
    cands = by_region.get(region, [])
    for st, units in cands:
        for u in units:
            if fold(u) == f:
                return st, u, 'name equals an existing Unit'
    for st, units in cands:
        if fold(st) == f:
            return st, None, 'name equals an existing Station'
    m = re.match(r'^(.*?)\s*(\d+)$', name.strip())
    if m:
        base = fold(m.group(1))
        for st, units in cands:
            if fold(st) == base:
                return st, None, f'Station name + number {m.group(2)}; that Unit does not exist yet'
    for other, olist in by_region.items():
        if other == region:
            continue
        for st, units in olist:
            for u in units:
                if fold(u) == f:
                    return st, u, f'exists in {other}, not {region}: the Region differs - check'
            if fold(st) == f:
                return st, None, f'exists in {other}, not {region}: the Region differs - check'
    best = max(cands, key=lambda c: difflib.SequenceMatcher(None, f, fold(c[0])).ratio(), default=None)
    if best and difflib.SequenceMatcher(None, f, fold(best[0])).ratio() >= 0.7:
        return best[0], None, 'similar spelling - check'
    return None, None, 'no similar Station - new Station?'

wb = Workbook()
ws = wb.active
ws.title = 'Names to review'
ws.sheet_view.rightToLeft = True
hdr = ['Region', 'Source name (as in files)', 'Storage vessels', 'Recovery tanks', 'Gas detectors', 'Hoses',
       'Installed SRVs', 'Total records', 'Suggested Station', 'Suggested Unit', 'Why suggested',
       'Your Station name', 'Your Unit name (leave empty if unknown)', 'Comment']
ws.append(hdr)
yellow = PatternFill('solid', fgColor='FFF2A8')
for c in ws[1]:
    c.font = Font(bold=True)
    c.alignment = Alignment(wrap_text=True, vertical='top')
kinds = {}
for region, name, sv, rt, gd, ho, srv in names:
    st, un, why = suggest(region, name)
    k = 'other Region' if why.startswith('exists in') else ('Station + number' if 'number' in why else why)
    kinds[k] = kinds.get(k, 0) + 1
    ws.append([region, name, sv, rt, gd, ho, srv, sv + rt + gd + ho + srv, st, un, why, None, None, None])
    for col in (12, 13):
        ws.cell(ws.max_row, col).fill = yellow
widths = [8, 34, 9, 9, 9, 7, 9, 9, 30, 26, 34, 30, 30, 24]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[ws.cell(1, i).column_letter].width = w
ws.freeze_panes = 'C2'
ws.auto_filter.ref = ws.dimensions

ref = wb.create_sheet('Existing Stations and Units')
ref.sheet_view.rightToLeft = True
ref.append(['Region', 'Station', 'Units'])
for c in ref[1]:
    c.font = Font(bold=True)
for region, st, units in stations:
    ref.append([region, st, ' ، '.join(units)])
for col, w in zip('ABC', (8, 36, 70)):
    ref.column_dimensions[col].width = w

how = wb.create_sheet('How to fill')
for line in [
    'Each row is one Station name, exactly as written in the source files, that matches no Station yet. Records are counted per type.',
    'Suggested Station / Unit are only SUGGESTIONS, and nothing is applied until you confirm it. The "Why suggested" column says how each was found.',
    'Fill the yellow columns: the Station it belongs to (copy the exact name from "Existing Stations and Units", or write a new Station name), and the Unit if known.',
    'To accept a suggestion as it is, write "ok" in "Your Station name".',
    'Leave the Unit empty when the source does not say which Unit.',
    'Leave the whole row empty to keep those records unlinked for now.',
]:
    how.append([line])
how.column_dimensions['A'].width = 140
wb.save(sys.argv[3])
print(len(names), 'names;', kinds)
